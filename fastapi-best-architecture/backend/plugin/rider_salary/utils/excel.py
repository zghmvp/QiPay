from __future__ import annotations

import contextlib
import csv
import io
import re

from datetime import date, datetime, timedelta
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from python_calamine import CalamineError, CalamineWorkbook

from backend.utils.timezone import timezone

HEADER_MAP: dict[str, str] = {
    '站点编码': 'site_code',
    '骑手工号': 'job_no',
    '订单号': 'order_no',
    '配送距离(公里)': 'distance_km',
    '商品重量(斤)': 'weight_jin',
    '下单时间': 'order_time',
    '送达时间': 'deliver_time',
    '订单状态': 'status',
    '订单金额': 'amount',
    '备注': 'remark',
}

REQUIRED_HEADERS: tuple[str, ...] = (
    '站点编码',
    '骑手工号',
    '订单号',
    '配送距离(公里)',
    '商品重量(斤)',
    '下单时间',
    '订单状态',
)

TEMPLATE_HEADERS: tuple[str, ...] = (
    '站点编码*',
    '骑手工号*',
    '订单号*',
    '配送距离(公里)*',
    '商品重量(斤)*',
    '下单时间*',
    '送达时间',
    '订单状态*',
    '订单金额',
    '备注',
)

_EXCEL_EXTS = {'.xlsx', '.xls'}
_CSV_EXTS = {'.csv'}
_FLEX_DT_RE = re.compile(r'^(\d{4})[-/](\d{1,2})[-/](\d{1,2})(?:\s+(\d{1,2}):(\d{2})(?::(\d{2}))?)?$')


class ExcelReadError(Exception):
    """Excel / CSV 读取错误"""

    def __init__(self, msg: str) -> None:
        self.msg = msg
        super().__init__(msg)


def normalize_header(name: object) -> str:
    """去掉星号与空白，全角括号转半角。"""
    text = '' if name is None else str(name).strip()
    text = text.replace('*', '').replace('＊', '')
    text = text.replace('（', '(').replace('）', ')')
    return text.replace(' ', '').replace('\u3000', '')


def parse_cell_datetime(value: object) -> datetime:
    """
    将单元格值解析为 FBA 时区感知 datetime

    :param value: 字符串、datetime、date 或 Excel 序列
    :return:
    """
    if value is None or (isinstance(value, str) and not value.strip()):
        raise ValueError('时间为空')
    if isinstance(value, datetime):
        return _ensure_tz(value)
    if isinstance(value, date) and not isinstance(value, datetime):
        return datetime(value.year, value.month, value.day, tzinfo=timezone.tz_info)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        epoch = datetime(1899, 12, 30, tzinfo=timezone.tz_info)
        return epoch + timedelta(days=float(value))
    text = str(value).strip()
    matched = _FLEX_DT_RE.match(text)
    if matched:
        year, month, day, hour, minute, second = matched.groups()
        return datetime(
            int(year),
            int(month),
            int(day),
            int(hour or 0),
            int(minute or 0),
            int(second or 0),
            tzinfo=timezone.tz_info,
        )
    raise ValueError(f'无法解析时间：{text}')


def read_rows(file_bytes: bytes, filename: str) -> list[dict[str, Any]]:
    """
    读取 xlsx/xls/csv，将中文表头映射为字段名

    :param file_bytes: 文件内容
    :param filename: 原始文件名（用于识别扩展名）
    :return:
    """
    ext = _file_ext(filename)
    if ext in _CSV_EXTS:
        raw_rows = _read_csv(file_bytes)
    elif ext in _EXCEL_EXTS:
        raw_rows = _read_excel(file_bytes)
    else:
        raise ExcelReadError('仅支持 Excel 或 CSV 文件')
    return _rows_to_dicts(raw_rows)


def write_workbook(sheets: list[tuple[str, list[str], list[list]]]) -> bytes:
    """
    将多个工作表写成 xlsx 字节

    :param sheets: (表名, 表头列表, 数据行列表)
    :return:
    """
    workbook = Workbook()
    default_sheet = workbook.active
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='305496')
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for index, (title, headers, rows) in enumerate(sheets):
        worksheet = default_sheet if index == 0 else workbook.create_sheet()
        worksheet.title = title[:31]
        for col, header in enumerate(headers, start=1):
            cell = worksheet.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
        for row_idx, row in enumerate(rows, start=2):
            for col, value in enumerate(row, start=1):
                worksheet.cell(row=row_idx, column=col, value=_excel_cell_value(value))
        for col, header in enumerate(headers, start=1):
            width = min(48, max(12, len(str(header)) + 4))
            worksheet.column_dimensions[get_column_letter(col)].width = width
        worksheet.freeze_panes = 'A2'
        worksheet.auto_filter.ref = worksheet.dimensions
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _ensure_tz(value: datetime) -> datetime:
    if value.tzinfo is None:
        return value.replace(tzinfo=timezone.tz_info)
    return timezone.from_datetime(value)


def _file_ext(filename: str) -> str:
    name = (filename or '').strip().rsplit('/', maxsplit=1)[-1]
    name = name.rsplit('\\', maxsplit=1)[-1]
    if '.' not in name:
        return ''
    return '.' + name.rsplit('.', maxsplit=1)[-1].lower()


def _read_excel(file_bytes: bytes) -> list[list[Any]]:
    try:
        workbook = CalamineWorkbook.from_filelike(io.BytesIO(file_bytes))
        sheet = workbook.get_sheet_by_index(0)
        rows = sheet.to_python(skip_empty_area=False)
        workbook.close()
    except (CalamineError, OSError, ValueError, IndexError) as exc:
        raise ExcelReadError('文件解析失败，请检查文件格式') from exc
    return rows


def _read_csv(file_bytes: bytes) -> list[list[Any]]:
    text = _decode_csv(file_bytes)
    reader = csv.reader(io.StringIO(text))
    return [list(row) for row in reader]


def _decode_csv(file_bytes: bytes) -> str:
    for encoding in ('utf-8-sig', 'utf-8', 'gbk'):
        with contextlib.suppress(UnicodeDecodeError):
            return file_bytes.decode(encoding)
    raise ExcelReadError('文件编码无法识别，请使用 UTF-8 或 GBK')


def _rows_to_dicts(raw_rows: list[list[Any]]) -> list[dict[str, Any]]:
    header_idx = None
    for idx, row in enumerate(raw_rows):
        if not _is_empty_row(row):
            header_idx = idx
            break
    if header_idx is None:
        raise ExcelReadError('文件为空，请按模板填写后重新上传')
    header_row = raw_rows[header_idx]
    field_by_col, missing = _map_headers(header_row)
    if missing:
        raise ExcelReadError(f'模板表头不符，缺少列：{"、".join(missing)}')
    result: list[dict[str, Any]] = []
    for offset, row in enumerate(raw_rows[header_idx + 1 :], start=header_idx + 2):
        if _is_empty_row(row):
            continue
        item: dict[str, Any] = {'_row': offset}
        for col, field in field_by_col.items():
            item[field] = row[col] if col < len(row) else None
        result.append(item)
    if not result:
        raise ExcelReadError('文件为空，请按模板填写后重新上传')
    return result


def _map_headers(header_row: list[Any]) -> tuple[dict[int, str], list[str]]:
    field_by_col: dict[int, str] = {}
    seen: set[str] = set()
    for col, cell in enumerate(header_row):
        normalized = normalize_header(cell)
        if not normalized:
            continue
        field = HEADER_MAP.get(normalized)
        if field is None:
            continue
        field_by_col[col] = field
        seen.add(normalized)
    missing = [name for name in REQUIRED_HEADERS if name not in seen]
    return field_by_col, missing


def _is_empty_row(row: list[Any]) -> bool:
    for cell in row:
        if cell is None:
            continue
        if isinstance(cell, str) and not cell.strip():
            continue
        return False
    return True


def _excel_cell_value(value: object) -> object:
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.replace(tzinfo=None)
    return value
