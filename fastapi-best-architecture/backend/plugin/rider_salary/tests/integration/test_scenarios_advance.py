"""E4、E10、E11、E15：预支额度、撤回、导出公式注入、需重算时的预估抵扣。"""

import io

from collections.abc import Generator
from decimal import Decimal

import pytest

from factories import (
    calculate_period,
    create_rider,
    create_site,
    expect_error,
    expect_ok,
    import_completed_orders,
    money,
    open_rider_account,
    order_row,
    provision_c01_month,
    use_shared_db_session,
)
from openpyxl import load_workbook
from runtime import ApiClient

_FORMULA_REASON = '=HYPERLINK("http://evil.example","点我")'


@pytest.fixture(autouse=True)
def _one_session_per_request() -> Generator[None, None, None]:
    """/me 写接口同时拿只读会话和事务会话，单连接上必须合成一个。"""
    with use_shared_db_session():
        yield


def _approve_and_pay(client: ApiClient, admin_token: dict[str, str], advance_id: int) -> None:
    expect_ok(
        client.post(
            f'/rider-salary/advances/{advance_id}/approve',
            headers=admin_token,
            json={'remark': '同意'},
        )
    )
    expect_ok(
        client.post(
            f'/rider-salary/advances/{advance_id}/mark-paid',
            headers=admin_token,
            json={'remark': '已发放'},
        )
    )


def test_e4_paid_outstanding_blocks_second_advance(client: ApiClient, admin_token: dict[str, str]) -> None:
    """E4 / P0-15：已发放 3000 且未抵扣时，额度剩余为 0，第二笔 3000 返回 400 且不建单。"""
    site = create_site(client, admin_token, name='额度站点')
    rider = create_rider(client, admin_token, site_id=site['id'], name='额度骑手', hire_date='2026-10-01')
    rider_headers = open_rider_account(client, admin_token, rider_id=rider['id'], job_no=rider['job_no'])
    first = expect_ok(
        client.post(
            '/rider-salary/me/advances',
            headers=rider_headers,
            json={'amount': '3000.00', 'reason': '第一次周转'},
        )
    )
    _approve_and_pay(client, admin_token, first['id'])
    quota = expect_ok(client.get('/rider-salary/me/advance-limit', headers=rider_headers))
    assert money(quota['available']) == Decimal('0.00')
    assert money(quota['outstanding_amount']) == money('3000.00')

    blocked = client.post(
        '/rider-salary/me/advances',
        headers=rider_headers,
        json={'amount': '3000.00', 'reason': '再次周转'},
    )
    message = expect_error(blocked)
    assert '额度' in message
    listed = expect_ok(
        client.get(
            '/rider-salary/advances',
            headers=admin_token,
            params={'rider_id': rider['id'], 'page': 1, 'size': 20},
        )
    )
    assert len(listed['items']) == 1
    assert listed['items'][0]['id'] == first['id']


def test_e10_rider_cannot_cancel_approved_advance(client: ApiClient, admin_token: dict[str, str]) -> None:
    """E10 / P0-15：骑手撤回待发放预支返回 400，单据仍是待发放。"""
    site = create_site(client, admin_token, name='撤回站点')
    rider = create_rider(client, admin_token, site_id=site['id'], name='撤回骑手', hire_date='2026-10-01')
    rider_headers = open_rider_account(client, admin_token, rider_id=rider['id'], job_no=rider['job_no'])
    created = expect_ok(
        client.post(
            '/rider-salary/me/advances',
            headers=rider_headers,
            json={'amount': '100.00', 'reason': '待发放撤回'},
        )
    )
    expect_ok(
        client.post(
            f'/rider-salary/advances/{created["id"]}/approve',
            headers=admin_token,
            json={'remark': '同意'},
        )
    )
    blocked = client.post(f'/rider-salary/me/advances/{created["id"]}/cancel', headers=rider_headers)
    message = expect_error(blocked)
    assert message
    current = expect_ok(client.get(f'/rider-salary/advances/{created["id"]}', headers=admin_token))
    assert current['status'] == 'to_pay'


def test_e11_advance_export_quotes_formula_reason(client: ApiClient, admin_token: dict[str, str]) -> None:
    """E11 / P3-07：以等号开头的预支原因导出后单元格类型是字符串，不是公式。"""
    site = create_site(client, admin_token, name='公式导出站点')
    rider = create_rider(client, admin_token, site_id=site['id'], name='公式导出骑手', hire_date='2026-10-01')
    rider_headers = open_rider_account(client, admin_token, rider_id=rider['id'], job_no=rider['job_no'])
    expect_ok(
        client.post(
            '/rider-salary/me/advances',
            headers=rider_headers,
            json={'amount': '20.00', 'reason': _FORMULA_REASON},
        )
    )
    exported = client.get(
        '/rider-salary/advances/export',
        headers=admin_token,
        params={'site_id': site['id']},
    )
    assert exported.status_code == 200, exported.text
    workbook = load_workbook(io.BytesIO(exported.content))
    worksheet = workbook.active
    assert worksheet is not None
    matched = [
        cell
        for row in worksheet.iter_rows()
        for cell in row
        if cell.value is not None and _FORMULA_REASON in str(cell.value)
    ]
    assert len(matched) == 1
    cell = matched[0]
    assert cell.data_type == 's'
    assert cell.column == 5
    assert str(cell.value).startswith("'")


def test_e15_stale_estimate_keeps_advance_deduction(client: ApiClient, admin_token: dict[str, str]) -> None:
    """E15 / P2-05：草稿抵扣 10 元后补一单，预估应发 20、抵扣 10、实发 10。"""
    ready = provision_c01_month(
        client,
        admin_token,
        month='2026-10',
        orders=[
            order_row('E15-01', '2026-10-01'),
            order_row('E15-02', '2026-10-02'),
            order_row('E15-03', '2026-10-03'),
        ],
        hire_date='2026-10-01',
        rider_name='预估骑手',
        site_name='预估站点',
    )
    rider = ready['rider']
    rider_headers = open_rider_account(client, admin_token, rider_id=rider['id'], job_no=rider['job_no'])
    created = expect_ok(
        client.post(
            '/rider-salary/me/advances',
            headers=rider_headers,
            json={'amount': '10.00', 'reason': '预估抵扣'},
        )
    )
    _approve_and_pay(client, admin_token, created['id'])
    calculate_period(client, admin_token, ready['period']['id'])
    fresh = expect_ok(client.get('/rider-salary/me/payroll-estimate', headers=rider_headers))
    assert money(fresh['gross']) == money('15.00')
    assert money(fresh['advance_deduction_estimate']) == money('10.00')
    assert money(fresh['net_estimate']) == money('5.00')
    assert fresh['is_estimate'] is False

    extra = import_completed_orders(
        client,
        admin_token,
        site_id=ready['site']['id'],
        site_code=ready['site']['code'],
        job_no=rider['job_no'],
        rows=[order_row('E15-04', '2026-10-04')],
    )
    assert extra['success_rows'] == 1
    stale = expect_ok(client.get('/rider-salary/me/payroll-estimate', headers=rider_headers))
    assert stale['is_estimate'] is True
    assert money(stale['gross']) == money('20.00')
    assert money(stale['advance_deduction_estimate']) == money('10.00')
    assert money(stale['net_estimate']) == money('10.00')
    kept = expect_ok(client.get(f'/rider-salary/advances/{created["id"]}', headers=admin_token))
    assert money(kept['remaining_amount']) == money('0.00')
    assert money(kept['deducted_amount']) == money('10.00')
