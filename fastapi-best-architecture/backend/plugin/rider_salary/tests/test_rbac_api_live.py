"""活 API 权限饱和目录（httpx 等价的 urllib 客户端）。"""

from __future__ import annotations

import io
import os

import pytest

from openpyxl import load_workbook

from backend.plugin.rider_salary.tests.rbac_live import (
    PREFIX,
    Actor,
    RbacWorld,
    _build_import_xlsx,
    api_reachable,
    assert_denied,
    assert_empty_list,
    bootstrap_world,
    denied,
    no_business_body,
    site_ids_of,
)

pytestmark = pytest.mark.integration


def _require_live() -> None:
    if api_reachable():
        return
    if os.environ.get('RIDER_SALARY_RBAC_LIVE') == '1':
        pytest.fail('RIDER_SALARY_RBAC_LIVE=1 但后端不可达')
    pytest.skip('后端未启动，跳过活 API 权限目录')


@pytest.fixture(scope='module')
def world() -> RbacWorld:
    _require_live()
    return bootstrap_world()


def _foreign_label(world: RbacWorld) -> str:
    return world.site_b.get('code') or world.site_b.get('name') or 'RBACB'


# --- 范围 / IDOR / 站跳 ---


def test_api_rbac_owner_a_lists_scoped(world: RbacWorld) -> None:
    """api-rbac-owner-a-lists-scoped · LIST-01/03/04"""
    ow = world.ow_a
    site_a = world.site_a_id
    for path in (f'{PREFIX}/orders', f'{PREFIX}/periods', f'{PREFIX}/adjustments', f'{PREFIX}/advances'):
        res = ow.call('GET', path, query={'page': 1, 'size': 50})
        assert res.status < 400, f'{path} {res.status} {res.msg}'
        leaked = site_ids_of(res.items) - {site_a}
        assert not leaked, f'{path} 列表含外站 {leaked}'
    sites = ow.call('GET', f'{PREFIX}/sites/all')
    ids = {int(row['id']) for row in (sites.data or []) if isinstance(row, dict)}
    assert ids == {site_a}, f'sites/all={ids}'
    dash = ow.call('GET', f'{PREFIX}/dashboard/summary')
    assert dash.status < 400, dash.msg
    text = str(dash.data)
    assert str(world.site_b_id) not in text or 'RBACB' not in text


def test_api_rbac_owner_a_idor_order_period_payroll(world: RbacWorld) -> None:
    """api-rbac-owner-a-idor-order-period-payroll · IDOR-01…07"""
    ow = world.ow_a
    assert_denied(ow.call('GET', f'{PREFIX}/orders/{world.order_b["id"]}'), label='IDOR-01 订单')
    assert_denied(
        ow.call('PUT', f'{PREFIX}/orders/{world.order_b["id"]}', json_body={'reason': '越权改单'}),
        label='IDOR-02 PUT',
    )
    assert_denied(
        ow.call('DELETE', f'{PREFIX}/orders/{world.order_b["id"]}', query={'reason': '越权删单'}),
        label='IDOR-02 DELETE',
    )
    if world.batch_b:
        assert_denied(ow.call('GET', f'{PREFIX}/import-batches/{world.batch_b["id"]}'), label='IDOR-03 批次')
        assert_denied(
            ow.call('GET', f'{PREFIX}/import-batches/{world.batch_b["id"]}/error-report'),
            label='IDOR-03 错误报告',
        )
    assert_denied(ow.call('GET', f'{PREFIX}/riders/{world.rider_b["id"]}'), label='IDOR-04 骑手')
    assert_denied(ow.call('GET', f'{PREFIX}/riders/{world.rider_b["id"]}/bindings'), label='IDOR-04 绑定')
    assert_denied(ow.call('GET', f'{PREFIX}/periods/{world.period_b["id"]}'), label='IDOR-05 周期')
    assert_denied(ow.call('GET', f'{PREFIX}/periods/{world.period_b["id"]}/export'), label='IDOR-05 导出')
    assert_denied(
        ow.call('POST', f'{PREFIX}/periods/{world.period_b["id"]}/calculate', json_body={'rider_ids': None}),
        label='IDOR-05 算薪',
    )
    assert_denied(
        ow.call('POST', f'{PREFIX}/periods/{world.period_b["id"]}/lock', json_body={'reason': '越权锁账'}),
        label='IDOR-05 锁账',
    )
    assert_denied(
        ow.call('POST', f'{PREFIX}/periods/{world.period_b["id"]}/mark-paid', json_body={'reason': '越权发薪'}),
        label='IDOR-05 发薪',
    )
    if world.payroll_b:
        assert_denied(ow.call('GET', f'{PREFIX}/payrolls/{world.payroll_b["id"]}'), label='IDOR-06 薪资单')
        assert_denied(ow.call('GET', f'{PREFIX}/payrolls/{world.payroll_b["id"]}/details'), label='IDOR-06 明细')
    assert_denied(ow.call('GET', f'{PREFIX}/calendar/{world.rider_b["id"]}'), label='IDOR-07 日历')
    assert_denied(ow.call('GET', f'{PREFIX}/calendar/{world.rider_b["id"]}/days/2026-07-15'), label='IDOR-07 日详情')


def test_api_rbac_owner_a_idor_advance_adjust_notice_site(world: RbacWorld) -> None:
    """api-rbac-owner-a-idor-advance-adjust-notice-site · IDOR-08…11"""
    ow = world.ow_a
    if world.advance_b:
        assert_denied(ow.call('GET', f'{PREFIX}/advances/{world.advance_b["id"]}'), label='IDOR-08 预支')
        assert_denied(
            ow.call('POST', f'{PREFIX}/advances/{world.advance_b["id"]}/approve', json_body={'remark': 'x'}),
            label='IDOR-08 审核',
        )
    assert_denied(ow.call('GET', f'{PREFIX}/adjustments/{world.adjustment_b["id"]}'), label='IDOR-09 奖惩')
    notice = ow.call('GET', f'{PREFIX}/notices/1')
    assert denied(notice) or notice.status == 404, '公告 API 已裁剪，残余不得 200 带体'
    assert_denied(ow.call('GET', f'{PREFIX}/sites/{world.site_b_id}'), label='IDOR-11 站点')
    assert_denied(
        ow.call(
            'PUT',
            f'{PREFIX}/sites/{world.site_b_id}/managers',
            json_body=[{'user_id': world.ow_a.user_id, 'role': 'owner'}],
        ),
        label='IDOR-11 managers',
    )


def test_api_rbac_owner_a_hop_query(world: RbacWorld) -> None:
    """api-rbac-owner-a-hop-query · HOP-01…06/16"""
    ow = world.ow_a
    site_b = world.site_b_id
    assert_denied(
        ow.call('GET', f'{PREFIX}/orders', query={'site_id': site_b, 'page': 1, 'size': 20}),
        label='HOP-01',
        allow_empty=True,
    )
    hop2 = ow.call('GET', f'{PREFIX}/orders', query={'rider_id': world.rider_b['id'], 'page': 1, 'size': 20})
    leaked = [row for row in hop2.items if int(row.get('site_id') or 0) == site_b]
    assert not leaked, 'HOP-02 带出站 B 订单'
    hop3 = ow.call('GET', f'{PREFIX}/periods', query={'site_id': site_b, 'page': 1, 'size': 20})
    assert_denied(hop3, label='HOP-03', allow_empty=False)
    hop4 = ow.call('GET', f'{PREFIX}/payrolls', query={'site_id': site_b, 'page': 1, 'size': 20})
    assert_empty_list(hop4, label='HOP-04')
    if world.period_b:
        hop4b = ow.call('GET', f'{PREFIX}/payrolls', query={'period_id': world.period_b['id'], 'page': 1, 'size': 20})
        assert_empty_list(hop4b, label='HOP-04 period')
    assert_denied(ow.call('GET', f'{PREFIX}/dashboard/summary', query={'site_id': site_b}), label='HOP-05')
    dash = ow.call('GET', f'{PREFIX}/dashboard/summary')
    assert dash.status < 400, dash.msg
    assert_denied(
        ow.call('GET', f'{PREFIX}/advances', query={'site_id': site_b, 'page': 1, 'size': 20}),
        label='HOP-16',
        allow_empty=False,
    )


def test_api_rbac_owner_a_hop_write_import_calc(world: RbacWorld) -> None:
    """api-rbac-owner-a-hop-write-import-calc · HOP-07…13"""
    ow = world.ow_a
    xlsx = _build_import_xlsx(world.site_b['code'], world.rider_b['job_no'], 'RBAC-HOP07-B')
    hop7 = ow.call(
        'POST',
        f'{PREFIX}/orders/import',
        form={'site_id': str(world.site_a_id), 'skip_errors': 'true', 'auto_recalc': 'false'},
        file_field=('file', 'hop7.xlsx', xlsx, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
    )
    if hop7.status < 400:
        data = hop7.data if isinstance(hop7.data, dict) else {}
        assert int(data.get('success_rows') or 0) == 0 or '无权' in str(data)
    hop8 = ow.call(
        'POST',
        f'{PREFIX}/orders',
        json_body={
            'order_no': 'RBAC-HOP08',
            'site_id': world.site_b_id,
            'rider_id': world.rider_b['id'],
            'distance_km': '1.00',
            'weight_jin': '1.00',
            'order_time': '2026-07-15T12:00:00+08:00',
            'status': 'completed',
        },
    )
    assert_denied(hop8, label='HOP-08')
    assert_denied(
        ow.call('POST', f'{PREFIX}/periods/generate', json_body={'site_id': world.site_b_id, 'month': '2026-06'}),
        label='HOP-09',
    )
    before = world.su.call('GET', f'{PREFIX}/periods/{world.period_b["id"]}')
    hop10 = ow.call(
        'POST',
        f'{PREFIX}/periods/{world.period_a["id"]}/calculate',
        json_body={'rider_ids': [world.rider_b['id']]},
    )
    assert hop10.status != 500
    after = world.su.call('GET', f'{PREFIX}/periods/{world.period_b["id"]}')
    assert (before.data or {}).get('status') == (after.data or {}).get('status')
    assert_denied(
        ow.call(
            'PUT',
            f'{PREFIX}/riders/{world.rider_a["id"]}',
            json_body={'site_id': world.site_b_id, 'reason': ' hop11'},
        ),
        label='HOP-11',
    )
    assert_denied(
        ow.call('PUT', f'{PREFIX}/riders/{world.rider_b["id"]}', json_body={'name': '改名', 'reason': 'hop12'}),
        label='HOP-12',
    )
    if world.plan_version:
        assert_denied(
            ow.call(
                'POST',
                f'{PREFIX}/riders/{world.rider_b["id"]}/bindings',
                json_body={
                    'plan_version_id': world.plan_version['id'],
                    'binding_type': 'default',
                    'start_date': '2026-07-01',
                    'end_date': None,
                    'reason': 'hop13',
                },
            ),
            label='HOP-13',
        )


def test_api_rbac_owner_b_mirror(world: RbacWorld) -> None:
    """api-rbac-owner-b-mirror · HOP-17"""
    ow = world.ow_b
    assert_denied(
        ow.call('GET', f'{PREFIX}/orders/{world.order_a["id"]}'),
        label='OW-B IDOR 订单 A',
    )
    assert_denied(
        ow.call('GET', f'{PREFIX}/periods', query={'site_id': world.site_a_id, 'page': 1, 'size': 10}),
        label='OW-B hop periods A',
    )
    sites = ow.call('GET', f'{PREFIX}/sites/all')
    ids = {int(row['id']) for row in (sites.data or []) if isinstance(row, dict)}
    assert ids == {world.site_b_id}


def test_api_rbac_empty_manager_sees_nothing(world: RbacWorld) -> None:
    """api-rbac-empty-manager-sees-nothing · HOP-14"""
    empty = world.empty
    sites = empty.call('GET', f'{PREFIX}/sites/all')
    assert sites.status < 500, sites.msg
    rows = sites.data if isinstance(sites.data, list) else sites.items
    assert rows == [] or sites.total == 0
    for path in (f'{PREFIX}/orders', f'{PREFIX}/periods', f'{PREFIX}/payrolls', f'{PREFIX}/riders'):
        res = empty.call('GET', path, query={'page': 1, 'size': 20})
        assert res.status != 500, f'{path} 空集 500：{res.msg}'
        assert_empty_list(res, label=f'EMPTY {path}')


def test_api_rbac_salary_admin_sees_all(world: RbacWorld) -> None:
    """api-rbac-salary-admin-sees-all · LIST-07"""
    orders = world.sa.call('GET', f'{PREFIX}/orders', query={'page': 1, 'size': 100})
    ids = site_ids_of(orders.items)
    assert world.site_a_id in ids and world.site_b_id in ids, f'SA 订单站点 {ids}'


def test_api_rbac_audit_site_scope(world: RbacWorld) -> None:
    """api-rbac-audit-site-scope · LIST-02 / 缺口 #3"""
    res = world.ow_a.call('GET', f'{PREFIX}/audit-logs', query={'page': 1, 'size': 100})
    assert res.status < 400, res.msg
    blob = str(res.data)
    leaked = [
        marker
        for marker in (_foreign_label(world), 'RBAC-OB-001', 'RBAC-ADJ-B', '权限站B')
        if marker and marker in blob
    ]
    assert not leaked, f'审计泄漏外站标记 {leaked}（缺口 #3）'


def test_api_rbac_so_global_catalog_read(world: RbacWorld) -> None:
    """api-rbac-so-global-catalog-read · LIST-08 / VERT-18"""
    plans = world.ow_a.call('GET', f'{PREFIX}/plans', query={'page': 1, 'size': 20})
    assert plans.status < 400, plans.msg
    versions = world.ow_a.call('GET', f'{PREFIX}/plan-versions', query={'page': 1, 'size': 20})
    assert versions.status < 400, versions.msg
    if world.plan_version:
        detail = world.ow_a.call('GET', f'{PREFIX}/plan-versions/{world.plan_version["id"]}')
        assert detail.status < 400, detail.msg
        text = str(detail.data)
        assert 'items' in text or 'formula' in text or 'formula_json' in text
        assert_denied(
            world.ow_a.call(
                'PUT',
                f'{PREFIX}/plan-versions/{world.plan_version["id"]}/items',
                json_body=[],
            ),
            label='SO PUT items',
        )
        assert_denied(
            world.ow_a.call(
                'POST',
                f'{PREFIX}/plan-versions/{world.plan_version["id"]}/trial',
                json_body={
                    'rider_id': world.rider_a['id'],
                    'start_date': '2026-07-01',
                    'end_date': '2026-07-31',
                    'mode': 'full_version',
                },
            ),
            label='SO trial',
        )


def test_api_rbac_so_sites_get_without_menu(world: RbacWorld) -> None:
    """api-rbac-so-sites-get-without-menu · VERT-17"""
    res = world.ow_a.call('GET', f'{PREFIX}/sites/{world.site_a_id}')
    assert res.status < 400, res.msg
    mgr = world.ow_a.call('GET', f'{PREFIX}/sites/{world.site_a_id}/managers')
    assert mgr.status < 400, mgr.msg


def test_api_rbac_read_hung_on_write_perms(world: RbacWorld) -> None:
    """api-rbac-read-hung-on-write-perms · LIST-09/10 与读挂写仍 200"""
    ow = world.ow_a
    riders = ow.call('GET', f'{PREFIX}/riders', query={'page': 1, 'size': 10})
    assert riders.status < 400, riders.msg
    rider = ow.call('GET', f'{PREFIX}/riders/{world.rider_a["id"]}')
    assert rider.status < 400, rider.msg
    adjs = ow.call('GET', f'{PREFIX}/adjustments', query={'page': 1, 'size': 10})
    assert adjs.status < 400, adjs.msg
    adv = ow.call('GET', f'{PREFIX}/advances', query={'page': 1, 'size': 10})
    assert adv.status < 400, adv.msg
    day = ow.call('GET', f'{PREFIX}/day-flags', query={'site_id': world.site_a_id, 'month': '2026-07'})
    assert day.status in {404, 405} or denied(day)
    notices = ow.call('GET', f'{PREFIX}/notices', query={'page': 1, 'size': 10})
    assert notices.status in {404, 405} or denied(notices) or notices.status < 400


# --- 垂直 ---


def test_api_rbac_owner_denied_site_account_plan_rollback(world: RbacWorld) -> None:
    """api-rbac-owner-denied-site-account-plan-rollback · VERT-01…04/11"""
    ow = world.ow_a
    assert_denied(
        ow.call('POST', f'{PREFIX}/sites', json_body={'code': 'HACK', 'name': 'x', 'settle_cycle': 'month'}),
        label='VERT-01 create site',
    )
    assert_denied(ow.call('DELETE', f'{PREFIX}/sites/{world.site_a_id}'), label='VERT-01 delete site')
    assert_denied(
        ow.call(
            'POST',
            f'{PREFIX}/riders/{world.rider_a["id"]}/open-account',
            json_body={'reason': '越权开户', 'password': 'Rider@123456'},
        ),
        label='VERT-02',
    )
    assert_denied(
        ow.call('POST', f'{PREFIX}/subjects', json_body={'code': 'X', 'name': 'x', 'direction': 'bonus'}),
        label='VERT-03 subject',
    )
    assert_denied(
        ow.call('POST', f'{PREFIX}/plans', json_body={'code': 'X', 'name': 'x', 'short_name': 'x', 'color': '#000'}),
        label='VERT-03 plan',
    )
    if world.plan_version:
        assert_denied(
            ow.call(
                'POST',
                f'{PREFIX}/plan-versions/{world.plan_version["id"]}/rollback',
                json_body={'reason': '越权回退'},
            ),
            label='VERT-04',
        )
        assert_denied(
            ow.call(
                'PUT',
                f'{PREFIX}/plan-versions/{world.plan_version["id"]}/items',
                json_body=[],
            ),
            label='VERT-11',
        )


def test_api_rbac_owner_denied_reverse(world: RbacWorld) -> None:
    """api-rbac-owner-denied-reverse · VERT-05"""
    assert_denied(
        world.ow_a.call(
            'POST',
            f'{PREFIX}/periods/{world.period_a["id"]}/reverse',
            json_body={'reason': '越权反冲'},
        ),
        label='VERT-05',
    )


def test_api_rbac_salary_admin_reverse_ok_rollback_denied(world: RbacWorld) -> None:
    """api-rbac-salary-admin-reverse-ok-rollback-denied · VERT-06/20"""
    if world.plan_version:
        rb = world.sa.call(
            'POST',
            f'{PREFIX}/plan-versions/{world.plan_version["id"]}/rollback',
            json_body={'reason': 'SA 回退'},
        )
        assert denied(rb), f'SA 回退应 403，实际 {rb.status} {rb.msg}'
        assert '超级管理员' in rb.msg or '权限' in rb.msg or rb.status == 403
    rev = world.sa.call(
        'POST',
        f'{PREFIX}/periods/{world.period_a["id"]}/reverse',
        json_body={'reason': 'SA 反冲目录探测'},
    )
    # 有码：不得因权限码 403；业务状态不足则 400/409
    assert '未分配菜单' not in rev.msg
    assert not (rev.status == 403 and '权限' in rev.msg and '超级' not in rev.msg)


def test_api_rbac_superuser_rollback_ok(world: RbacWorld) -> None:
    """api-rbac-superuser-rollback-ok · 用未使用 draft"""
    created = world.su.call(
        'POST',
        f'{PREFIX}/plans',
        json_body={
            'code': 'RBAC-ROLL',
            'name': '回退夹具方案',
            'short_name': '回退',
            'color': '#112233',
            'status': 'enable',
        },
    )
    plan = created.data if isinstance(created.data, dict) else None
    if not plan or not plan.get('id'):
        plans = world.su.call('GET', f'{PREFIX}/plans', query={'page': 1, 'size': 50})
        plan = next((row for row in plans.items if row.get('code') == 'RBAC-ROLL'), None)
    if not plan:
        pytest.skip(f'无法创建回退夹具方案 {created.status} {created.msg}')
    ver = world.su.call('POST', f'{PREFIX}/plan-versions', json_body={'plan_id': plan['id'], 'mode_tag': 'custom'})
    version = ver.data if isinstance(ver.data, dict) else None
    if not version or not version.get('id'):
        pytest.skip(f'无法创建 draft 版本 {ver.status} {ver.msg}')
    preview = world.su.call('GET', f'{PREFIX}/plan-versions/{version["id"]}/rollback-preview')
    assert preview.status < 400 or denied(preview) is False or preview.status == 200
    rb = world.su.call(
        'POST',
        f'{PREFIX}/plan-versions/{version["id"]}/rollback',
        json_body={'reason': '饱和目录回退夹具'},
    )
    assert rb.status != 403, f'SU 回退被拒 {rb.status} {rb.msg}'


def test_api_rbac_sa_trial_foreign_rider(world: RbacWorld) -> None:
    """api-rbac-sa-trial-foreign-rider · VERT-19"""
    if not world.plan_version:
        pytest.skip('无启用方案版本')
    res = world.sa.call(
        'POST',
        f'{PREFIX}/plan-versions/{world.plan_version["id"]}/trial',
        json_body={
            'rider_id': world.rider_b['id'],
            'start_date': '2026-07-01',
            'end_date': '2026-07-31',
            'mode': 'full_version',
        },
    )
    # 现实现：全站 + 不 assert 骑手站 → 200；不得 403 无权访问站点
    assert res.status != 403 or '无权访问该站点' not in res.msg
    world.notes.append(f'SA trial 外站骑手 HTTP {res.status} {res.msg}')


def test_api_rbac_deputy_equals_owner(world: RbacWorld) -> None:
    """api-rbac-deputy-equals-owner · VERT-12"""
    dp = world.dp_a
    assert_denied(dp.call('GET', f'{PREFIX}/orders/{world.order_b["id"]}'), label='副 IDOR 订单')
    assert_denied(dp.call('GET', f'{PREFIX}/riders/{world.rider_b["id"]}'), label='副 IDOR 骑手')
    assert_denied(
        dp.call('POST', f'{PREFIX}/periods/{world.period_a["id"]}/reverse', json_body={'reason': '副反冲'}),
        label='副反冲',
    )
    own = dp.call('GET', f'{PREFIX}/sites/{world.site_a_id}')
    assert own.status < 400, own.msg


def test_api_rbac_rider_denied_admin_writes(world: RbacWorld) -> None:
    """api-rbac-rider-denied-admin-writes · VERT-07/16"""
    rd = world.rd_b
    assert_denied(
        rd.call('POST', f'{PREFIX}/periods/{world.period_a["id"]}/calculate', json_body={'rider_ids': None}),
        label='RD calculate',
    )
    assert_denied(
        rd.call('POST', f'{PREFIX}/periods/{world.period_a["id"]}/lock', json_body={'reason': 'x'}),
        label='RD lock',
    )
    xlsx = _build_import_xlsx(world.site_b['code'], world.rider_b['job_no'], 'RD-IMP')
    assert_denied(
        rd.call(
            'POST',
            f'{PREFIX}/orders/import',
            form={'site_id': str(world.site_b_id), 'skip_errors': 'true'},
            file_field=('file', 'rd.xlsx', xlsx, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
        ),
        label='RD import',
    )


def test_api_rbac_rider_rbac_no_menu(world: RbacWorld) -> None:
    """api-rbac-rider-rbac-no-menu · VERT-08/15"""
    rd = world.rd_b
    for path in (f'{PREFIX}/periods', f'{PREFIX}/payrolls', f'{PREFIX}/sites'):
        res = rd.call('GET', path, query={'page': 1, 'size': 10})
        assert denied(res), f'{path} 应为无菜单 403，实际 {res.status} {res.msg}'
        assert '菜单' in res.msg or res.status == 403


def test_api_rbac_staff_denied_me(world: RbacWorld) -> None:
    """api-rbac-staff-denied-me · VERT-10"""
    for actor in (world.su, world.sa, world.ow_a):
        res = actor.call('POST', f'{PREFIX}/me/advances', json_body={'amount': '1.00', 'reason': 'staff'})
        assert denied(res), f'{actor.code} /me 应 403，实际 {res.status} {res.msg}'
        assert '骑手' in res.msg


# --- /me ---


def test_api_rbac_me_self_only(world: RbacWorld) -> None:
    """api-rbac-me-self-only"""
    rd = world.rd_b
    profile = rd.call('GET', f'{PREFIX}/me/profile')
    assert profile.status < 400, profile.msg
    blob = str(profile.data)
    assert world.rider_a['job_no'] not in blob
    cal = rd.call('GET', f'{PREFIX}/me/calendar', query={'month': '2026-07'})
    assert cal.status < 400, cal.msg
    assert 'calc_trace' not in str(cal.data)
    day = rd.call('GET', f'{PREFIX}/me/days/2026-07-15')
    assert day.status < 400 or denied(day)
    assert world.order_a['order_no'] not in str(day.data)
    for path in (
        f'{PREFIX}/me/payroll-estimate',
        f'{PREFIX}/me/adjustments',
        f'{PREFIX}/me/plan',
        f'{PREFIX}/me/advances',
    ):
        res = rd.call('GET', path)
        assert res.status < 400 or denied(res)
        assert world.rider_a['job_no'] not in str(res.data)
        assert 'RBAC-OA-001' not in str(res.data)


def test_api_rbac_me_cancel_foreign_advance(world: RbacWorld) -> None:
    """api-rbac-me-cancel-foreign-advance · IDOR-12"""
    if not world.advance_a:
        pytest.skip('无 RD-A pending 预支')
    res = world.rd_b.call('POST', f'{PREFIX}/me/advances/{world.advance_a["id"]}/cancel', json_body={})
    assert res.status == 404 or res.code == 404
    assert '不存在' in res.msg


def test_api_rbac_me_notices_scope(world: RbacWorld) -> None:
    """api-rbac-me-notices-scope"""
    res = world.rd_b.call('GET', f'{PREFIX}/me/notices')
    if res.status == 404:
        pytest.skip('骑手公告已随产品裁剪')
    assert res.status < 400, res.msg
    for row in res.items or (res.data or []):
        if not isinstance(row, dict):
            continue
        site_id = row.get('site_id')
        assert site_id in {None, world.site_b_id}


def test_api_rbac_resigned_rider_forbidden(world: RbacWorld) -> None:
    """api-rbac-resigned-rider-forbidden"""
    from backend.plugin.rider_salary.tests.rbac_live import swagger_login

    data = swagger_login(world.resigned_username, os.environ.get('RIDER_SALARY_STAFF_PASS', 'Rider@123456'))
    actor = Actor('RX', world.resigned_username, '', data['access_token'])
    res = actor.call('GET', f'{PREFIX}/me/profile')
    assert denied(res)
    assert '不是有效骑手' in res.msg


# --- 注入 / 导出 / JWT-only / 正路径 ---


def test_api_rbac_query_injection(world: RbacWorld) -> None:
    """api-rbac-query-injection · QINJ-01…06"""
    ow = world.ow_a
    q1 = ow.call('GET', f'{PREFIX}/orders', query={'site_id': '1 OR 1=1', 'page': 1, 'size': 10})
    assert q1.status in {400, 422} or q1.total == 0 or denied(q1)
    q1b = ow.call('GET', f'{PREFIX}/orders', query={'site_id': -1, 'page': 1, 'size': 10})
    assert q1b.status != 200 or not q1b.items or denied(q1b)
    q2 = ow.call('GET', f'{PREFIX}/orders', query={'order_no': "' OR 1=1 --", 'page': 1, 'size': 10})
    assert q2.status < 500
    leaked = [row for row in q2.items if int(row.get('site_id') or 0) == world.site_b_id]
    assert not leaked
    q3 = ow.call('GET', f'{PREFIX}/periods', query={'rider_id': 0, 'page': 1, 'size': 50})
    assert not (site_ids_of(q3.items) - {world.site_a_id})
    q4 = ow.call('GET', f'{PREFIX}/orders', query={'page': -1, 'size': 99999})
    assert q4.status != 500
    q5 = ow.call('GET', f'{PREFIX}/periods', query={'month': '2026-13'})
    assert q5.status in {400, 422} or denied(q5) or q5.total == 0
    xlsx = _build_import_xlsx(world.site_b['code'], world.rider_b['job_no'], 'QINJ06')
    q6 = ow.call(
        'POST',
        f'{PREFIX}/orders/import',
        form={'site_id': str(world.site_a_id), 'skip_errors': 'true'},
        file_field=('file', 'qinj6.xlsx', xlsx, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
    )
    if q6.status < 400 and isinstance(q6.data, dict):
        assert int(q6.data.get('success_rows') or 0) == 0 or '无权' in str(q6.data)


def test_api_rbac_export_orders_period_advance(world: RbacWorld) -> None:
    """api-rbac-export-orders-period-advance · EXP-01…07"""
    ow = world.ow_a
    exp_b = ow.call('GET', f'{PREFIX}/orders/export', query={'site_id': world.site_b_id})
    assert_denied(exp_b, label='EXP-01')
    exp_a = ow.call('GET', f'{PREFIX}/orders/export', query={'site_id': world.site_a_id})
    assert exp_a.status < 400, exp_a.msg
    if exp_a.raw:
        wb = load_workbook(io.BytesIO(exp_a.raw))
        values = ' '.join(str(c.value) for row in wb.active.iter_rows() for c in row)
        assert world.site_b['code'] not in values
        disp = exp_a.headers.get('content-disposition', '')
        assert '订单' in disp or 'xlsx' in disp or disp
    assert_denied(ow.call('GET', f'{PREFIX}/periods/{world.period_b["id"]}/export'), label='EXP-03')
    adv = ow.call('GET', f'{PREFIX}/advances/export')
    if adv.status < 400 and adv.raw.startswith(b'PK'):
        wb = load_workbook(io.BytesIO(adv.raw))
        values = ' '.join(str(c.value) for row in wb.active.iter_rows() for c in row)
        assert 'RBACB' not in values and '权限站B' not in values
    assert_denied(ow.call('GET', f'{PREFIX}/advances/export', query={'site_id': world.site_b_id}), label='EXP-05')
    if world.batch_b:
        assert_denied(ow.call('GET', f'{PREFIX}/import-batches/{world.batch_b["id"]}/error-report'), label='EXP-06')


def test_api_rbac_rider_jwt_orders_batches(world: RbacWorld) -> None:
    """api-rbac-rider-jwt-orders-batches · 补 DependsRBAC 后须 403，且无业务行"""
    rd = world.rd_b
    for path in (f'{PREFIX}/orders', f'{PREFIX}/import-batches'):
        res = rd.call('GET', path, query={'page': 1, 'size': 20})
        assert denied(res), f'{path} 骑手 JWT 未 403：{res.status} {res.msg}'
        assert no_business_body(res)
        assert 'RBAC-OA' not in str(res.data) and 'RBAC-OB' not in str(res.data)
    detail = rd.call('GET', f'{PREFIX}/orders/{world.order_b["id"]}')
    assert denied(detail) and no_business_body(detail)
    if world.batch_b:
        batch = rd.call('GET', f'{PREFIX}/import-batches/{world.batch_b["id"]}')
        assert denied(batch) and no_business_body(batch)


def test_api_rbac_rider_jwt_engine(world: RbacWorld) -> None:
    """api-rbac-rider-jwt-engine · 补 DependsRBAC 后须 403"""
    rd = world.rd_b
    fields = rd.call('GET', f'{PREFIX}/engine/fields')
    assert denied(fields), f'RD engine fields {fields.status} {fields.msg}'
    val = rd.call(
        'POST',
        f'{PREFIX}/engine/validate',
        json_body={'stage': 'per_order', 'formula_json': {'kind': 'literal', 'value': 1}},
    )
    assert denied(val)
    ev = rd.call(
        'POST',
        f'{PREFIX}/engine/evaluate-sample',
        json_body={
            'stage': 'per_order',
            'formula_json': {'kind': 'literal', 'value': 1},
            'context': {'单量': 1},
        },
    )
    assert denied(ev)
    assert 'RBAC-RA' not in str(ev.data)


def test_api_rbac_calc_lock_import_same_site(world: RbacWorld) -> None:
    """api-rbac-calc-lock-import-same-site · 本站正路径不得被过滤写反"""
    ow = world.ow_a
    calc = ow.call(
        'POST',
        f'{PREFIX}/periods/{world.period_a["id"]}/calculate',
        json_body={'rider_ids': None},
    )
    assert calc.status != 403, f'本站算薪被 403：{calc.msg}'
    lock = ow.call(
        'POST',
        f'{PREFIX}/periods/{world.period_a["id"]}/lock',
        json_body={'reason': '饱和目录本站锁账探测'},
    )
    # 业务硬拦（无方案/未算出）可以 400，但不能是「无权访问该站点」
    assert '无权访问该站点' not in lock.msg
    xlsx = _build_import_xlsx(world.site_a['code'], world.rider_a['job_no'], 'RBAC-OWN-IMP')
    imp = ow.call(
        'POST',
        f'{PREFIX}/orders/import',
        form={'site_id': str(world.site_a_id), 'skip_errors': 'true', 'auto_recalc': 'false'},
        file_field=('file', 'own.xlsx', xlsx, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
    )
    assert imp.status != 403, f'本站导入 403 {imp.msg}'
