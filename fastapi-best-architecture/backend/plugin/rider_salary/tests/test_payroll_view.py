"""P0-08 有效薪资单口径。

P1-03 集成基座（``tests/integration/``）尚未就绪，验收写成服务层单测：
E13 同一订单只保留有效单上的一行；回退作废的单不进导出；
周期应发、实发等于有效单合计。Q-06 按推荐方案 A：净差 = 当前有效补发 − 原单，
并另给最终实发。多轮反冲补发只认最新一张未被反冲的补发。
"""

from datetime import date
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

from openpyxl import load_workbook

from backend.plugin.rider_salary.enums import PayrollKind, PayrollStatus
from backend.plugin.rider_salary.service.export_service import (
    HISTORY_HEADERS,
    NET_HEADERS,
    SHEET_DETAIL,
    SHEET_NET,
    SHEET_ORIGINAL,
    SHEET_REVERSAL,
    SHEET_SUMMARY,
    SHEET_SUPPLEMENT,
    build_payroll_export_sheets,
)
from backend.plugin.rider_salary.service.payroll_view import (
    build_rider_views,
    details_of_effective,
    pick_effective_payroll,
    sum_effective_gross,
    summarize_period_stats,
)
from backend.plugin.rider_salary.utils.excel import write_workbook


def _slip(**kwargs: object) -> SimpleNamespace:
    data: dict[str, object] = {
        'id': 1,
        'period_id': 9,
        'rider_id': 2,
        'kind': PayrollKind.normal.value,
        'status': PayrollStatus.finalized.value,
        'deleted': 0,
        'reversed': False,
        'reversed_of_id': None,
        'stale': False,
        'calc_version': 1,
        'net': Decimal('0.00'),
        'gross': Decimal('0.00'),
        'order_count': 1,
        'valid_order_count': 1,
        'per_order_total': Decimal('0.00'),
        'daily_total': Decimal('0.00'),
        'period_total': Decimal('0.00'),
        'bonus_total': Decimal('0.00'),
        'penalty_total': Decimal('0.00'),
        'deduction_total': Decimal('0.00'),
        'advance_deduction': Decimal('0.00'),
        'calc_time': None,
    }
    data.update(kwargs)
    return SimpleNamespace(**data)


def _detail(payroll_id: int, amount: str, *, detail_id: int, order_id: int = 70) -> SimpleNamespace:
    return SimpleNamespace(
        id=detail_id,
        payroll_id=payroll_id,
        rider_id=2,
        subject_id=3,
        amount=Decimal(amount),
        stage='per_order',
        include_in_gross=True,
        source='formula',
        biz_date=date(2026, 9, 2),
        plan_version_id=None,
        order_id=order_id,
        calc_trace=None,
        deleted=0,
    )


def _chain() -> list[SimpleNamespace]:
    """原单 → 反冲 → 补发 → 再反冲 → 再补发，外加一张作废草稿。"""
    return [
        _slip(id=1, net=Decimal('100.00'), gross=Decimal('100.00'), reversed=True),
        _slip(
            id=2,
            kind=PayrollKind.reversal.value,
            net=Decimal('-100.00'),
            gross=Decimal('-100.00'),
            reversed_of_id=1,
            calc_version=1,
        ),
        _slip(
            id=3,
            kind=PayrollKind.supplement.value,
            net=Decimal('90.00'),
            gross=Decimal('90.00'),
            reversed=True,
            calc_version=2,
        ),
        _slip(
            id=4,
            kind=PayrollKind.reversal.value,
            net=Decimal('-90.00'),
            gross=Decimal('-90.00'),
            reversed_of_id=3,
            calc_version=2,
        ),
        _slip(
            id=5,
            kind=PayrollKind.supplement.value,
            net=Decimal('80.00'),
            gross=Decimal('80.00'),
            calc_version=3,
        ),
        _slip(id=6, status=PayrollStatus.voided.value, net=Decimal('50.00'), gross=Decimal('50.00')),
    ]


def _e13_details() -> list[SimpleNamespace]:
    return [
        _detail(1, '4.00', detail_id=11),
        _detail(2, '-4.00', detail_id=12),
        _detail(3, '4.00', detail_id=13),
        _detail(4, '-4.00', detail_id=14),
        _detail(5, '4.00', detail_id=15),
    ]


def test_pick_keeps_latest_unreversed_supplement_across_rounds() -> None:
    chosen = pick_effective_payroll(_chain())
    assert chosen is not None
    assert chosen.id == 5


def test_pick_is_empty_after_reverse_before_next_supplement() -> None:
    payrolls = [row for row in _chain() if row.id != 5]
    assert pick_effective_payroll(payrolls) is None


def test_reversed_supplement_does_not_hide_a_live_original() -> None:
    payrolls = [
        _slip(id=1, net=Decimal('100.00'), calc_version=1),
        _slip(
            id=2,
            kind=PayrollKind.supplement.value,
            net=Decimal('90.00'),
            reversed=True,
            calc_version=4,
        ),
    ]
    chosen = pick_effective_payroll(payrolls)
    assert chosen is not None
    assert chosen.id == 1


def test_stale_draft_is_still_the_effective_slip() -> None:
    chosen = pick_effective_payroll([
        _slip(id=1, reversed=True),
        _slip(
            id=2,
            kind=PayrollKind.supplement.value,
            status=PayrollStatus.draft.value,
            stale=True,
            net=Decimal('80.00'),
            calc_version=2,
        ),
    ])
    assert chosen is not None
    assert chosen.id == 2


def test_e13_order_keeps_one_detail_from_the_effective_slip() -> None:
    kept = details_of_effective(_e13_details(), _chain())
    assert [row.amount for row in kept] == [Decimal('4.00')]
    assert [row.payroll_id for row in kept] == [5]


def test_e13_drops_every_line_when_the_latest_supplement_is_reversed() -> None:
    payrolls = [row for row in _chain() if row.id != 5]
    assert details_of_effective(_e13_details(), payrolls) == []


def test_net_diff_is_latest_supplement_minus_original() -> None:
    view = build_rider_views(_chain(), _e13_details())[0]
    assert view.original_net == Decimal('100.00')
    assert view.reversal_net == Decimal('-100.00')
    assert view.supplement_net == Decimal('80.00')
    assert view.net_diff == Decimal('-20.00')
    assert view.final_net == Decimal('80.00')
    assert [row.id for row in view.details] == [15]


def test_net_diff_is_zero_when_there_is_no_supplement() -> None:
    view = build_rider_views([_slip(id=1, net=Decimal('100.00'), gross=Decimal('100.00'))])[0]
    assert view.net_diff == Decimal('0.00')
    assert view.final_net == Decimal('100.00')
    assert view.supplement_net == Decimal('0.00')


def test_net_diff_is_zero_while_waiting_for_the_next_supplement() -> None:
    payrolls = [row for row in _chain() if row.id != 5]
    view = build_rider_views(payrolls)[0]
    assert view.effective is None
    assert view.net_diff == Decimal('0.00')
    assert view.final_net == Decimal('0.00')
    assert view.original_net == Decimal('100.00')
    assert view.has_live is True


def test_period_totals_equal_effective_slips_and_skip_voided() -> None:
    other = _slip(id=7, rider_id=8, net=Decimal('20.00'), gross=Decimal('20.00'))
    stats = summarize_period_stats([*_chain(), other])[9]
    assert stats['gross_total'] == Decimal('100.00')
    assert stats['net_total'] == Decimal('100.00')
    assert stats['payroll_count'] == 6
    assert stats['rider_count'] == 2
    assert stats['kind_counts'] == {'normal': 2, 'reversal': 2, 'supplement': 2}
    assert sum_effective_gross([*_chain(), other]) == Decimal('100.00')


def test_voided_only_period_contributes_nothing() -> None:
    stats = summarize_period_stats([
        _slip(id=6, status=PayrollStatus.voided.value, stale=True, net=Decimal('50.00'), gross=Decimal('50.00')),
    ])[9]
    assert stats['payroll_count'] == 0
    assert stats['rider_count'] == 0
    assert stats['stale_count'] == 0
    assert stats['gross_total'] == Decimal('0.00')
    assert stats['net_total'] == Decimal('0.00')


def test_export_sheets_follow_effective_view_and_keep_history() -> None:
    riders = {2: SimpleNamespace(job_no='R002', name='李明')}
    sheets = build_payroll_export_sheets(
        _chain(),
        _e13_details(),
        riders,
        site_name='东站',
        period_text='2026-09-01~2026-09-30',
        plan_names={},
        subject_names={3: '基础单价'},
        order_nos={70: 'ORD-1'},
    )
    names = [name for name, _headers, _rows in sheets]
    assert names == [SHEET_SUMMARY, SHEET_DETAIL, SHEET_NET, SHEET_ORIGINAL, SHEET_REVERSAL, SHEET_SUPPLEMENT]
    content = write_workbook(sheets)
    workbook = load_workbook(BytesIO(content))

    summary = workbook[SHEET_SUMMARY]
    assert [cell.value for cell in summary[1]][:6] == ['工号', '姓名', '站点', '周期', '单据类型', '状态']
    assert summary.max_row == 2
    assert summary['A2'].value == 'R002'
    assert summary['E2'].value == '补发'
    assert summary['Q2'].value == 80

    detail = workbook[SHEET_DETAIL]
    assert detail.max_row == 2
    assert detail['G2'].value == 'ORD-1'
    assert detail['H2'].value == 4

    net = workbook[SHEET_NET]
    assert [cell.value for cell in net[1]] == NET_HEADERS
    assert [cell.value for cell in net[2]] == ['R002', '李明', 100, -100, 80, -20, 80]

    original = workbook[SHEET_ORIGINAL]
    assert [cell.value for cell in original[1]] == HISTORY_HEADERS
    assert original.max_row == 2
    assert original['Q2'].value == 100
    assert original['S2'].value == '是'

    reversal = workbook[SHEET_REVERSAL]
    assert reversal.max_row == 3
    assert [reversal['Q2'].value, reversal['Q3'].value] == [-100, -90]
    assert reversal['T2'].value == 1
    assert reversal['T3'].value == 3

    supplement = workbook[SHEET_SUPPLEMENT]
    assert [supplement['Q2'].value, supplement['Q3'].value] == [90, 80]
    assert [supplement['S2'].value, supplement['S3'].value] == ['是', '否']
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(min_row=2, max_col=20, values_only=True):
            assert 50 not in row
            assert 50.0 not in row
