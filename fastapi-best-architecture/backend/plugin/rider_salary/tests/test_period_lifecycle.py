from datetime import date
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

import pytest

from openpyxl import load_workbook

from backend.common.exception import errors
from backend.plugin.rider_salary.enums import CycleType, PayrollKind, PayrollStatus, PeriodStatus
from backend.plugin.rider_salary.service.export_service import (
    ADJUSTMENT_HEADERS,
    DETAIL_HEADERS,
    NET_HEADERS,
    SHEET_ADJUSTMENT,
    SHEET_DETAIL,
    SHEET_NET,
    SHEET_SUMMARY,
    SUMMARY_HEADERS,
    summarize_trace,
)
from backend.plugin.rider_salary.service.period_service import (
    assert_can_transition,
    covering_period_ranges,
    reversible_payrolls,
    stale_lock_message,
)
from backend.plugin.rider_salary.utils.excel import write_workbook


def test_generate_half_month_two_periods() -> None:
    ranges = covering_period_ranges(CycleType.half_month, None, 2026, 9)
    assert ranges == [
        (date(2026, 9, 1), date(2026, 9, 15)),
        (date(2026, 9, 16), date(2026, 9, 30)),
    ]


def test_generate_month_one_period() -> None:
    ranges = covering_period_ranges(CycleType.month, None, 2026, 9)
    assert ranges == [(date(2026, 9, 1), date(2026, 9, 30))]


def test_generate_custom_anchor_covers_month() -> None:
    ranges = covering_period_ranges(CycleType.custom, {'anchor_day': 26}, 2026, 9)
    assert ranges == [
        (date(2026, 8, 26), date(2026, 9, 25)),
        (date(2026, 9, 26), date(2026, 10, 25)),
    ]


def test_lock_rejects_stale_drafts() -> None:
    msg = stale_lock_message(['A001', 'A002'])
    assert msg == '存在需重算的薪资结果：工号 A001、A002，请先重算'


def test_illegal_transition_rejected() -> None:
    with pytest.raises(errors.RequestError, match='结算周期当前状态为开放，不允许执行标记发薪'):
        assert_can_transition(PeriodStatus.open.value, PeriodStatus.paid.value)
    with pytest.raises(errors.RequestError, match='结算周期当前状态为已发薪，不允许执行锁账'):
        assert_can_transition(PeriodStatus.paid.value, PeriodStatus.locked.value)
    assert assert_can_transition(PeriodStatus.open.value, PeriodStatus.locked.value) == '锁账'
    assert assert_can_transition(PeriodStatus.reopened.value, PeriodStatus.locked.value) == '锁账'
    assert assert_can_transition(PeriodStatus.locked.value, PeriodStatus.paid.value) == '标记发薪'
    assert assert_can_transition(PeriodStatus.locked.value, PeriodStatus.reopened.value) == '反冲补发'
    assert assert_can_transition(PeriodStatus.paid.value, PeriodStatus.reopened.value) == '反冲补发'


def test_reverse_selects_finalized_and_skips_reversal() -> None:
    payrolls = [
        SimpleNamespace(
            kind=PayrollKind.normal.value,
            status=PayrollStatus.finalized.value,
            reversed=False,
            net=Decimal('100.00'),
        ),
        SimpleNamespace(
            kind=PayrollKind.reversal.value,
            status=PayrollStatus.finalized.value,
            reversed=False,
            net=Decimal('-100.00'),
        ),
        SimpleNamespace(
            kind=PayrollKind.normal.value,
            status=PayrollStatus.draft.value,
            reversed=False,
            net=Decimal('8.00'),
        ),
        SimpleNamespace(
            kind=PayrollKind.supplement.value,
            status=PayrollStatus.paid.value,
            reversed=True,
            net=Decimal('12.00'),
        ),
        SimpleNamespace(
            kind=PayrollKind.normal.value,
            status=PayrollStatus.paid.value,
            reversed=False,
            net=Decimal('50.00'),
        ),
    ]
    selected = reversible_payrolls(payrolls)
    assert [item.net for item in selected] == [Decimal('100.00'), Decimal('50.00')]


def test_export_workbook_headers() -> None:
    content = write_workbook([
        (SHEET_SUMMARY, SUMMARY_HEADERS, [['A001', '张伟']]),
        (SHEET_DETAIL, DETAIL_HEADERS, []),
        (SHEET_ADJUSTMENT, ADJUSTMENT_HEADERS, []),
        (SHEET_NET, NET_HEADERS, [['A001', '张伟', 100, -100, 90, 90]]),
    ])
    workbook = load_workbook(BytesIO(content))
    assert workbook.sheetnames == [SHEET_SUMMARY, SHEET_DETAIL, SHEET_ADJUSTMENT, SHEET_NET]
    assert [cell.value for cell in workbook[SHEET_SUMMARY][1]] == SUMMARY_HEADERS
    assert [cell.value for cell in workbook[SHEET_DETAIL][1]] == DETAIL_HEADERS
    assert [cell.value for cell in workbook[SHEET_ADJUSTMENT][1]] == ADJUSTMENT_HEADERS
    assert [cell.value for cell in workbook[SHEET_NET][1]] == NET_HEADERS
    assert workbook[SHEET_SUMMARY]['A2'].value == 'A001'


def test_summarize_trace_one_line_chinese() -> None:
    text = summarize_trace({
        '条件': '配送距离 > 5',
        '条件结果': True,
        '公式': '(配送距离 - 5) * 0.8',
        '变量': {'配送距离': 7.2},
        '结果': 1.76,
    })
    assert '条件：配送距离 > 5（真）' in text
    assert '公式：(配送距离 - 5) * 0.8' in text
    assert '配送距离=7.2' in text
    assert '结果=1.76' in text


def test_c5_current_round_payable_uses_latest_unreversed_supplement() -> None:
    from types import SimpleNamespace

    from backend.plugin.rider_salary.service.export_service import _current_round_payable

    def row(pk, kind, net, *, version=1, reversed_=False, reversed_of=None):
        return SimpleNamespace(
            id=pk,
            rider_id=1,
            kind=kind,
            net=Decimal(net),
            calc_version=version,
            reversed=reversed_,
            reversed_of_id=reversed_of,
        )

    # 原单 100 被反冲，第一轮补发 90 又被反冲，第二轮补发 95
    payrolls = [
        row(1, 'normal', '100', version=1, reversed_=True),
        row(2, 'reversal', '-100', version=1, reversed_of=1),
        row(3, 'supplement', '90', version=2, reversed_=True),
        row(4, 'reversal', '-90', version=2, reversed_of=3),
        row(5, 'supplement', '95', version=3),
    ]
    # 本轮应付 = 冲掉上一轮的 -90 + 本轮补发 95
    assert _current_round_payable(payrolls) == Decimal('5')
    # 没有补发时退回本期净差
    assert _current_round_payable(payrolls[:2]) == Decimal('0')
