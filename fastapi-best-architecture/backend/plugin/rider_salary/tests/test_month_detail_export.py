"""P6-06：自然月明细只导出一个文件，半月结用周期列区分，行不先收成列表。"""

from collections.abc import AsyncIterator
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import anyio
import pytest

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet._write_only import WriteOnlyWorksheet
from sqlalchemy import select

from backend.common.exception import errors
from backend.plugin.rider_salary.enums import CalcStage, DetailSource
from backend.plugin.rider_salary.model.payroll_detail import RiderSalaryPayrollDetail
from backend.plugin.rider_salary.service.month_detail_export import (
    MONTH_DETAIL_HEADERS,
    MONTH_DETAIL_SHEET,
    _detail_filters,
    _iter_month_detail_rows,
    build_month_detail_row,
    month_detail_export_service,
    month_detail_included,
    parse_export_month,
    payroll_ids_for_month_export,
    pick_period_for_day,
)
from backend.plugin.rider_salary.utils.excel import EXPORT_ROW_LIMIT_MSG, StreamingWorkbook, append_streamed_rows
from backend.utils.timezone import timezone

_PERIOD_COLUMN = MONTH_DETAIL_HEADERS.index('周期')


def _period(pk: int, start: date, end: date, status: str, *, rider_id: int = 0) -> SimpleNamespace:
    return SimpleNamespace(
        id=pk,
        rider_id=rider_id,
        site_id=2,
        start_date=start,
        end_date=end,
        status=status,
    )


def _moment(day: int) -> datetime:
    return datetime(2026, 9, day, 8, 0, tzinfo=timezone.tz_info)


def _detail_tuple(day: int, period_id: int, order_no: str) -> tuple[object, ...]:
    return (
        date(2026, 9, day),
        'D5A001',
        '张伟',
        order_no,
        Decimal('1.20'),
        Decimal('2.00'),
        _moment(day),
        None,
        'completed',
        '底薪A',
        3,
        '配送费',
        Decimal('5.00'),
        'formula',
        True,
        period_id,
    )


def _order_tuple(day: int, order_no: str) -> tuple[object, ...]:
    return (
        date(2026, 9, day),
        7,
        'D5A001',
        '张伟',
        order_no,
        Decimal('3.00'),
        Decimal('1.00'),
        _moment(day),
        None,
        'cancelled',
    )


def _half_month_periods() -> list[SimpleNamespace]:
    return [
        _period(11, date(2026, 9, 1), date(2026, 9, 15), 'locked'),
        _period(12, date(2026, 9, 16), date(2026, 9, 30), 'open'),
    ]


def test_month_rows_keep_orders_daily_formulas_and_manual_subjects() -> None:
    assert month_detail_included(source=DetailSource.formula.value, stage=CalcStage.per_order.value) is True
    assert month_detail_included(source=DetailSource.formula.value, stage=CalcStage.daily.value) is True
    assert month_detail_included(source=DetailSource.manual.value, stage=CalcStage.daily.value) is True
    assert month_detail_included(source=DetailSource.formula.value, stage=CalcStage.period.value) is False
    assert month_detail_included(source=DetailSource.advance.value, stage=CalcStage.period.value) is False
    assert month_detail_included(source=DetailSource.reversal.value, stage=CalcStage.per_order.value) is False


def test_reversed_payroll_is_left_out_of_the_month_file() -> None:
    payrolls = [
        SimpleNamespace(
            id=1,
            period_id=11,
            rider_id=7,
            deleted=0,
            status='finalized',
            kind='normal',
            reversed=True,
            calc_version=1,
        ),
        SimpleNamespace(
            id=2,
            period_id=11,
            rider_id=7,
            deleted=0,
            status='finalized',
            kind='reversal',
            reversed=False,
            calc_version=1,
            reversed_of_id=1,
        ),
        SimpleNamespace(
            id=3,
            period_id=11,
            rider_id=7,
            deleted=0,
            status='finalized',
            kind='supplement',
            reversed=False,
            calc_version=2,
        ),
    ]
    assert payroll_ids_for_month_export(payrolls) == [3]


def test_parse_september_and_reject_bad_month() -> None:
    text, start, end = parse_export_month('2026-09')
    assert text == '2026-09'
    assert start == date(2026, 9, 1)
    assert end == date(2026, 9, 30)
    with pytest.raises(errors.RequestError) as exc:
        parse_export_month('2026-9')
    assert exc.value.msg == '月份格式应为 YYYY-MM'


def test_rider_period_wins_over_site_period() -> None:
    site = _period(11, date(2026, 9, 1), date(2026, 9, 15), 'locked')
    rider = _period(21, date(2026, 9, 1), date(2026, 9, 15), 'open', rider_id=7)
    picked = pick_period_for_day([site, rider], 7, date(2026, 9, 3))
    assert picked is rider
    second = pick_period_for_day(_half_month_periods(), 7, date(2026, 9, 20))
    assert second is not None
    assert second.id == 12


def test_detail_sql_keeps_order_daily_and_manual_inside_the_month() -> None:
    stmt = select(RiderSalaryPayrollDetail.id).where(*_detail_filters([3], date(2026, 9, 1), date(2026, 9, 30), 7))
    compiled = str(stmt.compile(compile_kwargs={'literal_binds': True}))
    assert 'manual' in compiled
    assert 'per_order' in compiled
    assert 'daily' in compiled
    assert '2026-09-01' in compiled
    assert '2026-09-30' in compiled


def test_half_month_september_is_one_file_with_period_column() -> None:
    """半月结的 9 月只得到 1 个文件，上半月和下半月都在周期列里。"""
    state = {'details': 0, 'orders': 0}
    periods = _half_month_periods()

    async def _fake_stream(db: object, stmt: object, yield_per: int = 500) -> AsyncIterator[tuple[object, ...]]:
        del db, stmt, yield_per
        if state['orders'] == 0 and state['details'] < 2:
            for day, period_id, order_no in ((3, 11, 'ORD-A'), (8, 11, 'ORD-A2')):
                state['details'] += 1
                await anyio.sleep(0)
                yield _detail_tuple(day, period_id, order_no)
            return
        state['orders'] += 1
        await anyio.sleep(0)
        yield _order_tuple(20, 'ORD-B')

    seen: list[tuple[int, int]] = []
    original = WriteOnlyWorksheet.append

    def _spy(self: WriteOnlyWorksheet, row: object) -> None:
        seen.append((state['details'], state['orders']))
        original(self, row)

    async def _run() -> bytes:
        book = StreamingWorkbook()
        sheet = book.add_sheet(MONTH_DETAIL_SHEET, MONTH_DETAIL_HEADERS)
        with patch(
            'backend.plugin.rider_salary.service.month_detail_export.stream_rows',
            _fake_stream,
        ):
            await append_streamed_rows(
                book,
                sheet,
                _iter_month_detail_rows(
                    SimpleNamespace(),
                    site_name='甲站',
                    periods=periods,
                    effective_ids=[3],
                    site_id=2,
                    rider_id=7,
                    start=date(2026, 9, 1),
                    end=date(2026, 9, 30),
                ),
            )
        return book.to_bytes()

    with patch.object(WriteOnlyWorksheet, 'append', _spy):
        content = anyio.run(_run)

    assert seen[0] == (0, 0)
    assert seen[1] == (1, 0)
    assert state['orders'] == 1
    workbook = load_workbook(BytesIO(content))
    assert workbook.sheetnames == [MONTH_DETAIL_SHEET]
    rows = list(workbook.active.iter_rows(values_only=True))
    assert rows[0][_PERIOD_COLUMN] == '周期'
    assert rows[0][MONTH_DETAIL_HEADERS.index('周期状态')] == '周期状态'
    periods_in_file = {row[_PERIOD_COLUMN] for row in rows[1:]}
    assert periods_in_file == {'2026-09-01~2026-09-15', '2026-09-16~2026-09-30'}
    statuses = {row[MONTH_DETAIL_HEADERS.index('周期状态')] for row in rows[1:]}
    assert statuses == {'已锁账', '开放'}
    assert rows[1][MONTH_DETAIL_HEADERS.index('科目')] == '配送费'
    assert rows[1][MONTH_DETAIL_HEADERS.index('来源')] == '公式'
    assert rows[1][MONTH_DETAIL_HEADERS.index('版本号')] == 'v3'
    assert rows[-1][MONTH_DETAIL_HEADERS.index('订单号')] == 'ORD-B'
    assert rows[-1][MONTH_DETAIL_HEADERS.index('科目')] in ('', None)


def test_export_rejects_over_fifty_thousand_before_reading_rows() -> None:
    async def _run() -> None:
        with (
            patch(
                'backend.plugin.rider_salary.service.month_detail_export.get_visible_site_ids',
                new=AsyncMock(return_value=None),
            ),
            patch(
                'backend.plugin.rider_salary.service.month_detail_export._get_site',
                new=AsyncMock(return_value=SimpleNamespace(id=2, name='甲站')),
            ),
            patch(
                'backend.plugin.rider_salary.service.month_detail_export._load_periods',
                new=AsyncMock(return_value=[]),
            ),
            patch(
                'backend.plugin.rider_salary.service.month_detail_export._load_payrolls',
                new=AsyncMock(return_value=[]),
            ),
            patch(
                'backend.plugin.rider_salary.service.month_detail_export._count',
                new=AsyncMock(return_value=50_001),
            ),
            patch(
                'backend.plugin.rider_salary.service.month_detail_export.stream_rows',
                side_effect=AssertionError('超限后不应读取明细'),
            ),
            patch(
                'backend.plugin.rider_salary.service.month_detail_export._iter_month_detail_rows',
                side_effect=AssertionError('超限后不应开始写行'),
            ),
        ):
            with pytest.raises(errors.RequestError) as exc:
                await month_detail_export_service.export(
                    db=SimpleNamespace(),
                    request=SimpleNamespace(),
                    site_id=2,
                    month='2026-09',
                    rider_id=None,
                )
        assert exc.value.msg == EXPORT_ROW_LIMIT_MSG

    anyio.run(_run)


def test_export_month_writes_one_workbook() -> None:
    flags: list[object] = []
    original = Workbook.__init__

    def _wrapped(self: Workbook, *args: object, **kwargs: object) -> None:
        flags.append(kwargs.get('write_only'))
        original(self, *args, **kwargs)

    async def _rows(*args: object, **kwargs: object) -> AsyncIterator[list[object]]:
        del args, kwargs
        await anyio.sleep(0)
        yield build_month_detail_row(
            biz_date=date(2026, 9, 3),
            job_no='D5A001',
            rider_name='张伟',
            site_name='甲站',
            order_no='ORD-A',
            subject_name='配送费',
            amount=Decimal('5.00'),
            source='formula',
            include_in_gross=True,
            period_text='2026-09-01~2026-09-15',
            period_status='已锁账',
        )
        yield build_month_detail_row(
            biz_date=date(2026, 9, 20),
            job_no='D5A001',
            rider_name='张伟',
            site_name='甲站',
            order_no='ORD-B',
            period_text='2026-09-16~2026-09-30',
            period_status='开放',
        )

    async def _run() -> tuple[bytes, str]:
        with (
            patch(
                'backend.plugin.rider_salary.service.month_detail_export.get_visible_site_ids',
                new=AsyncMock(return_value=None),
            ),
            patch(
                'backend.plugin.rider_salary.service.month_detail_export._get_site',
                new=AsyncMock(return_value=SimpleNamespace(id=2, name='甲站')),
            ),
            patch(
                'backend.plugin.rider_salary.service.month_detail_export._get_rider',
                new=AsyncMock(return_value=SimpleNamespace(id=7, name='张伟')),
            ),
            patch(
                'backend.plugin.rider_salary.service.month_detail_export._load_periods',
                new=AsyncMock(return_value=_half_month_periods()),
            ),
            patch(
                'backend.plugin.rider_salary.service.month_detail_export._load_payrolls',
                new=AsyncMock(return_value=[]),
            ),
            patch(
                'backend.plugin.rider_salary.service.month_detail_export._count',
                new=AsyncMock(return_value=2),
            ),
            patch(
                'backend.plugin.rider_salary.service.month_detail_export._iter_month_detail_rows',
                _rows,
            ),
            patch(
                'backend.plugin.rider_salary.service.month_detail_export.audit_service.record',
                new=AsyncMock(),
            ),
            patch.object(Workbook, '__init__', _wrapped),
        ):
            return await month_detail_export_service.export(
                db=SimpleNamespace(),
                request=SimpleNamespace(),
                site_id=2,
                month='2026-09',
                rider_id=7,
            )

    content, filename = anyio.run(_run)
    assert flags == [True]
    assert filename == '当月明细_甲站_张伟_2026-09.xlsx'
    assert content[:2] == b'PK'
    workbook = load_workbook(BytesIO(content))
    assert len(workbook.sheetnames) == 1
    body = list(workbook.active.iter_rows(values_only=True))
    assert {row[_PERIOD_COLUMN] for row in body[1:]} == {
        '2026-09-01~2026-09-15',
        '2026-09-16~2026-09-30',
    }
